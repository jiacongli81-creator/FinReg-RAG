import re
import uuid
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

from ..config import settings
from ..db import insert_chunks, insert_document, utc_now
from .embeddings import embed


SUPPORTED = {".pdf", ".txt", ".docx", ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}
DATE_PATTERN = re.compile(r"(?:20\d{2})[年./-](?:\d{1,2})[月./-](?:\d{1,2})日?")


def _normalize(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pages(path: Path) -> list[tuple[int, str]]:
    if path.suffix.lower() == ".txt":
        return [(1, path.read_text(encoding="utf-8", errors="ignore"))]
    if path.suffix.lower() == ".docx":
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return [(1, "\n".join(parts))]
    if path.suffix.lower() in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff"}:
        try:
            import pytesseract
            from PIL import Image
        except ImportError as exc:
            raise ValueError("图片需要安装 OCR 依赖 pytesseract 和 Tesseract 引擎") from exc
        return [(1, pytesseract.image_to_string(Image.open(path), lang="chi_sim+eng"))]
    reader = PdfReader(str(path))
    return [(index + 1, page.extract_text() or "") for index, page in enumerate(reader.pages)]


def _chunk_pages(pages: list[tuple[int, str]]) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    for page, raw_text in pages:
        text = _normalize(raw_text)
        if not text:
            continue
        start = 0
        while start < len(text):
            end = min(len(text), start + settings.chunk_size)
            if end < len(text):
                boundary = max(text.rfind("\n", start, end), text.rfind("。", start, end))
                if boundary > start + settings.chunk_size // 2:
                    end = boundary + 1
            content = text[start:end].strip()
            if content:
                chunks.append({"id": str(uuid.uuid4()), "content": content, "page": page, "section": None})
            if end >= len(text):
                break
            start = max(end - settings.chunk_overlap, start + 1)
    return chunks


def ingest_file(path: Path) -> dict[str, Any]:
    if path.suffix.lower() not in SUPPORTED:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    document_id = str(uuid.uuid4())
    title = path.stem
    date_match = DATE_PATTERN.search(path.name)
    effective_date = date_match.group(0) if date_match else None
    validity = "unknown"
    if any(word in path.name for word in ("失效", "废止", "作废")):
        validity = "invalid"
    try:
        chunks = _chunk_pages(_extract_pages(path))
        if not chunks:
            raise ValueError("没有提取到文字；这可能是扫描版 PDF，需要安装 OCR 引擎")
        for chunk in chunks:
            chunk.update(document_id=document_id, embedding=embed(chunk["content"]))
        insert_document(
            {
                "id": document_id,
                "title": title,
                "filename": path.name,
                "file_type": path.suffix.lower().lstrip("."),
                "effective_date": effective_date,
                "validity": validity,
                "created_at": utc_now(),
                "status": "ready",
                "error": None,
            }
        )
        insert_chunks(chunks)
        return {"id": document_id, "title": title, "filename": path.name, "file_type": path.suffix.lower().lstrip("."), "chunk_count": len(chunks), "created_at": utc_now(), "status": "ready", "error": None}
    except Exception as exc:
        insert_document(
            {
                "id": document_id,
                "title": title,
                "filename": path.name,
                "file_type": path.suffix.lower().lstrip("."),
                "effective_date": effective_date,
                "validity": validity,
                "created_at": utc_now(),
                "status": "failed",
                "error": str(exc),
            }
        )
        raise
